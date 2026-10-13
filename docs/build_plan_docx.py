from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Plan_Analisis_Predictivo_ERP_Guia_7_Fases.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="000000", size=9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_row(row):
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "17365D")
        set_cell_borders(cell)
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=9)
        if widths:
            cell.width = Inches(widths[index])
    mark_header_row(table.rows[0])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_borders(cells[index])
            set_cell_text(cells[index], str(value), size=9)
            if row_index % 2:
                set_cell_shading(cells[index], "F2F6FA")
            if widths:
                cells[index].width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_page_field(paragraph):
    run = paragraph.add_run("Página ")
    run.font.size = Pt(8)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name != "Title" else "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")

footer = section.footer
footer_paragraph = footer.paragraphs[0]
footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer_paragraph.add_run("Sistema JulMarp ERP · Plan de análisis predictivo · ")
add_page_field(footer_paragraph)

title = doc.add_paragraph(style="Title")
title.add_run("Plan de análisis predictivo para el ERP")
title.paragraph_format.space_after = Pt(4)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
run = subtitle.add_run("Hoja de ruta para producción, inventario, calidad, mantenimiento, finanzas y planeación")
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(70, 90, 115)

doc.add_paragraph(
    "Este documento define cómo evolucionar el ERP hacia un sistema que no solo registre operaciones, "
    "sino que anticipe riesgos, explique sus causas y recomiende acciones. El trabajo comenzará con la "
    "base común de datos y trazabilidad, continuará con MFG y después se extenderá al resto de los departamentos."
)

doc.add_heading("Resumen ejecutivo", level=1)
doc.add_paragraph(
    "El sistema ya cuenta con un módulo predictivo inicial en MFG: reportes KPI, simulación de periodos, "
    "pronóstico de indicadores, rangos de incertidumbre, comparación de tendencias y alertas. La siguiente "
    "etapa es convertir ese prototipo en una capacidad transversal, medible y reutilizable."
)
add_table(
    doc,
    ["Área", "Situación actual", "Siguiente resultado"],
    [
        ("MFG / Producción", "Pronóstico de KPI y alertas", "Predicción por línea, producto y turno"),
        ("Inventario y Compras", "Datos de entradas, salidas y proveedores", "Demanda, reorden y riesgo de ruptura"),
        ("Calidad y QA", "Defectos, scrap e inspecciones", "Riesgo de defecto y causa probable"),
        ("Mantenimiento", "Datos indirectos de recursos y paros", "Bitácora y riesgo de falla"),
        ("Finanzas", "Costos, presupuestos, pagos y cobros", "Flujo, desviación y rentabilidad"),
    ],
    widths=[1.45, 2.55, 2.55],
)

doc.add_heading("Arquitectura común", level=1)
doc.add_paragraph(
    "Todos los departamentos utilizarán el mismo ciclo de trabajo. La diferencia estará en sus métricas, "
    "fuentes y acciones, no en la forma de registrar y evaluar las predicciones."
)
doc.add_paragraph(
    "Datos operativos → Dataset histórico → Modelo → Pronóstico o riesgo → Acción → Resultado real",
    style="Subtitle",
)
add_table(
    doc,
    ["Capa", "Responsabilidad", "Resultado"],
    [
        ("Datos", "Capturar operaciones y eventos", "Historial confiable"),
        ("Preparación", "Limpiar, agrupar y validar", "Dataset utilizable"),
        ("Modelo", "Detectar patrones y estimar futuro", "Predicción"),
        ("Explicación", "Mostrar causas, confianza y rango", "Interpretación"),
        ("Acción", "Crear alerta, tarea o recomendación", "Decisión operativa"),
        ("Evaluación", "Comparar predicción contra resultado", "Error y mejora"),
    ],
    widths=[1.15, 3.0, 2.4],
)

doc.add_heading("Estado actual del proyecto", level=1)
add_bullets(
    doc,
    [
        "MFG cuenta con pronóstico de OEE, disponibilidad, rendimiento, calidad y producción.",
        "La interfaz permite comparar métricas, activar o desactivar series y observar tendencias en Base 100.",
        "La simulación local puede generar 10 periodos semanales para alimentar el análisis.",
        "La base local SQLite está activa para desarrollo y Supabase permanece desactivado.",
        "Se agregó RegistroPrediccion para guardar modelo, versión, periodo, confianza, rango, acción y resultado futuro.",
        "La migración 0031_registroprediccion ya fue aplicada en la base local.",
    ],
)

doc.add_heading("Plan de trabajo por fases", level=1)

doc.add_heading("Fase 1 Base común y calidad de datos", level=2)
add_numbered(
    doc,
    [
        "Crear el catálogo común de métricas, fuentes, unidades, objetivos y umbrales.",
        "Normalizar periodos, líneas, departamentos, turnos y responsables.",
        "Registrar muestras, nulos, duplicados y cobertura temporal.",
        "Guardar cada predicción con modelo, versión, valor, rango y confianza.",
        "Comparar automáticamente cada predicción con el resultado real.",
    ],
)

doc.add_heading("Fase 2 MFG y Producción", level=2)
add_bullets(
    doc,
    [
        "Separar predicciones por línea, producto y turno.",
        "Agregar scrap, retraso de órdenes y costo final.",
        "Comparar tendencia lineal, promedio móvil y modelos alternativos.",
        "Convertir alertas en tareas operativas.",
        "No mostrar alta confianza con menos de cuatro periodos.",
    ],
)

doc.add_heading("Fase 3 Inventario y Compras", level=2)
add_bullets(
    doc,
    [
        "Demanda futura por material.",
        "Punto de reorden y fecha probable de agotamiento.",
        "Riesgo de ruptura de stock.",
        "Tiempo de entrega por proveedor.",
        "Cantidad sugerida de compra.",
    ],
)
doc.add_paragraph("Datos base: entradas, salidas, transferencias, inventario, proveedores, precios y órdenes de compra.")

doc.add_heading("Fase 4 Calidad y QA", level=2)
add_bullets(
    doc,
    [
        "Probabilidad de defecto por producto, BOM, línea o proveedor.",
        "Tendencia de scrap y reincidencia de defectos.",
        "Riesgo de reclamo de cliente.",
        "Proceso o material con mayor contribución al problema.",
        "Acción correctiva sugerida y seguimiento del resultado.",
    ],
)

doc.add_heading("Fase 5 Mantenimiento y recursos", level=2)
doc.add_paragraph(
    "El ERP todavía no cuenta con un módulo formal de mantenimiento. La primera versión se apoyará en "
    "paros, horas de máquina, costos, cambios de ciclo y defectos relacionados con recursos."
)
add_bullets(
    doc,
    [
        "Crear bitácora de mantenimiento.",
        "Registrar fallas, causas, duración y reparaciones.",
        "Calcular riesgo de falla por máquina.",
        "Integrar sensores de temperatura, vibración o consumo cuando estén disponibles.",
    ],
)

doc.add_heading("Fase 6 Finanzas y Planeación", level=2)
add_bullets(
    doc,
    [
        "Flujo de caja proyectado.",
        "Cobros y pagos atrasados.",
        "Desviación presupuestal.",
        "Costo final de órdenes.",
        "Rentabilidad estimada.",
        "Capacidad futura contra demanda.",
    ],
)

doc.add_heading("Fase 7 Operación y gobierno de modelos", level=2)
add_bullets(
    doc,
    [
        "Reentrenamiento programado.",
        "Versionado de modelos.",
        "Monitoreo de error y deriva de datos.",
        "Registro de predicciones y decisiones.",
        "Permisos por departamento.",
        "Respaldos y migración a PostgreSQL para producción.",
        "Pruebas automatizadas de datos, modelos y vistas.",
    ],
)

doc.add_heading("Prioridad inmediata", level=1)
doc.add_paragraph("El trabajo inmediato seguirá este orden:")
add_numbered(
    doc,
    [
        "Completar el contrato común de predicciones y sus métricas de calidad.",
        "Registrar los resultados reales y calcular el error posterior.",
        "Migrar completamente MFG al registro histórico.",
        "Construir el primer caso predictivo de Inventario y Compras.",
        "Agregar alertas accionables y permisos por departamento.",
    ],
)

doc.add_heading("Criterios de éxito", level=1)
add_bullets(
    doc,
    [
        "Cada predicción identifica departamento, métrica, periodo, modelo y versión.",
        "El usuario ve valor actual, valor esperado, confianza, rango y acción.",
        "El sistema mide el error cuando existe un resultado real.",
        "Cada departamento cuenta con al menos un caso predictivo útil.",
        "Las alertas pueden convertirse en tareas o seguimientos dentro del ERP.",
        "Los modelos pueden compararse, actualizarse y auditarse.",
    ],
)

doc.add_heading("Cómo continuar el trabajo", level=1)
doc.add_paragraph(
    "Para retomar el proyecto, comparte este documento y señala la fase o actividad que deseas continuar. "
    "La prioridad actual es terminar la Fase 1 y ampliar el análisis desde MFG hacia Inventario y Compras."
)

doc.add_page_break()
doc.add_heading("Guía manual de revisión de las siete fases", level=1)
doc.add_paragraph(
    "Esta guía anexa al plan explica cómo revisar manualmente los ajustes realizados en las siete fases, "
    "qué debe observar el usuario y cómo comprobar que el ciclo predictivo funciona con la base local. "
    "Está dirigida a usuarios funcionales, responsables de cada departamento y personal técnico que necesite "
    "validar una ejecución completa."
)

doc.add_heading("Antes de comenzar", level=2)
add_bullets(
    doc,
    [
        "Trabajar con el proyecto local en C:\\1_Proyectos\\Sistema-JulMarp-ERP.",
        "Confirmar que la base activa es SQLite local y que la conexión de Supabase permanece desactivada, sin eliminar su configuración.",
        "Usar el entorno virtual del proyecto para ejecutar Django.",
        "Ingresar con un usuario con permisos de consulta para todos los módulos o con el usuario administrador de pruebas.",
    ],
)
doc.add_paragraph("Comprobación técnica inicial:")
for command in [
    "venv\\Scripts\\python.exe manage.py check",
    "venv\\Scripts\\python.exe manage.py makemigrations --check --dry-run",
    "venv\\Scripts\\python.exe manage.py migrate --plan",
]:
    paragraph = doc.add_paragraph(style="Intense Quote")
    run = paragraph.add_run(command)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(9.5)
doc.add_paragraph(
    "Los tres comandos deben terminar sin errores. El segundo debe indicar que no hay cambios de migración pendientes "
    "y el tercero no debe mostrar operaciones por ejecutar."
)

doc.add_heading("Mapa rápido de las fases", level=2)
add_table(
    doc,
    ["Fase", "Qué se ajustó", "Cómo se comprueba"],
    [
        ("1 Datos y contrato", "Catálogo común, calidad, versiones y registro de predicciones", "Admin y ciclo predictivo"),
        ("2 MFG", "Pronóstico de KPI, segmentos, alertas y comparación de tendencias", "Indicadores y predicción de producción"),
        ("3 Inventario", "Demanda, reorden, faltantes y riesgo de inventario", "Almacén y órdenes de compra"),
        ("4 Calidad", "Riesgo de defecto, scrap y segmentos de calidad", "SQA, OQA y defectos"),
        ("5 Mantenimiento", "Paros, riesgo de falla y uso de recursos", "Control de recursos MFG"),
        ("6 Finanzas", "Costos, presupuesto, flujo y desviaciones", "Dashboard financiero"),
        ("7 Operación", "Ciclo integrado, evaluación, salud, deriva y trazabilidad", "Comando de ciclo y auditoría"),
    ],
    widths=[0.8, 3.1, 2.65],
)

doc.add_heading("Fase 1 Datos y contrato predictivo", level=2)
doc.add_paragraph(
    "La primera fase concentra la estructura que usan todas las áreas. Cada predicción debe quedar identificada "
    "por departamento, métrica, periodo, valor estimado, rango de incertidumbre, nivel de confianza, modelo, versión "
    "y acción sugerida. Cuando ya existe el resultado real, el sistema calcula el error y conserva la trazabilidad."
)
add_numbered(
    doc,
    [
        "Abrir /admin/ y revisar CatalogoMetricaPredictiva. Confirme que existan métricas activas para MFG, Inventario, Calidad, Mantenimiento y Finanzas.",
        "Revisar RegistroPrediccion. Verifique que cada fila tenga periodo objetivo, departamento, métrica, versión y estado de evaluación.",
        "Confirme que las versiones 1.0 y 1.1 estén diferenciadas. La versión 1.1 corresponde al ciclo operativo más reciente.",
        "Si el registro muestra predicciones pendientes, no es necesariamente un error: significa que todavía falta capturar el resultado real del periodo objetivo.",
    ],
)

doc.add_heading("Fase 2 MFG y Producción", level=2)
doc.add_paragraph(
    "MFG es el módulo con mayor cobertura visual. El análisis usa producción, reportes KPI, tiempos, scrap, calidad y "
    "recursos para proyectar indicadores. La pantalla permite activar o desactivar métricas y comparar una o varias series "
    "sin obligar al usuario a mostrar todo al mismo tiempo."
)
add_numbered(
    doc,
    [
        "Abrir /auth/produccion/indicadores-kpis/. Revise que el periodo, la producción total, el scrap y los KPI tengan datos.",
        "Abrir /auth/produccion/prediccion-produccion/. Seleccione una métrica y después active las métricas adicionales desde la lista de comparación.",
        "Compruebe que cada serie tenga color y leyenda propios. Use Base 100 cuando las unidades sean distintas; use la vista normal cuando necesite comparar porcentajes directos.",
        "Revise el rango de incertidumbre y la alerta principal. Una alerta debe llevar una recomendación breve, no solo un valor numérico.",
        "Abrir /auth/produccion/control-recursos/. Registre o consulte horas de máquina, paros, ciclos y causas para que el análisis de recursos tenga contexto.",
    ],
)

doc.add_heading("Fase 3 Inventario y Compras", level=2)
doc.add_paragraph(
    "Esta fase transforma movimientos de almacén y compras en señales operativas: demanda esperada, punto de reorden, "
    "riesgo de faltante, fecha probable de agotamiento y proveedor que requiere atención."
)
add_numbered(
    doc,
    [
        "Abrir /auth/inventario/almacen/. Revise existencias, entradas, salidas, consumo y materiales con riesgo.",
        "Abrir /auth/inventario/ordenes-compra/. Compare las órdenes abiertas con la demanda esperada y el tiempo de entrega.",
        "Confirme que una recomendación de compra tenga material, cantidad, fecha o prioridad y proveedor relacionado.",
        "Si la calidad del dato aparece baja, revise primero movimientos faltantes, periodos sin consumo y unidades inconsistentes.",
    ],
)

doc.add_heading("Fase 4 Calidad y QA", level=2)
doc.add_paragraph(
    "La fase de Calidad usa defectos, scrap, inspecciones y reclamos para anticipar dónde puede repetirse un problema. "
    "La revisión debe centrarse en la causa y en la acción correctiva, no solo en el porcentaje de defecto."
)
add_numbered(
    doc,
    [
        "Abrir /auth/qa/sqa/, /auth/qa/oqa/, /auth/qa/qqaat-defectos/ y /auth/qa/customer-service/.",
        "Verifique que los defectos estén relacionados con producto, proceso, línea, proveedor o periodo.",
        "Compare la tendencia histórica con la proyección y revise los segmentos con mayor riesgo.",
        "Documente la acción correctiva y vuelva a revisar el indicador cuando exista un nuevo resultado real.",
    ],
)

doc.add_heading("Fase 5 Mantenimiento y recursos", level=2)
doc.add_paragraph(
    "La fase 5 agrega una bitácora formal de paros de mantenimiento y conecta los recursos de producción con el riesgo "
    "de falla. En la simulación actual pueden existir predicciones de mantenimiento aun cuando no se hayan cargado paros; "
    "para mejorar su calidad se deben registrar fallas y reparaciones reales."
)
add_numbered(
    doc,
    [
        "Abrir /auth/produccion/control-recursos/. Revise máquinas, horas de uso, cambios de ciclo y causas de paro.",
        "En el administrador, revisar RegistroParoMantenimiento y capturar máquina, fecha, duración, causa, reparación y costo cuando ocurra un evento.",
        "Compare el riesgo de falla por máquina con la bitácora. Un riesgo alto sin eventos registrados debe investigarse como falta de datos.",
        "Confirme que el resultado de la acción preventiva pueda quedar documentado para evaluar el pronóstico después.",
    ],
)

doc.add_heading("Fase 6 Finanzas y Planeación", level=2)
doc.add_paragraph(
    "La fase financiera consolida costos, presupuesto, pagos, cobros, flujo y rentabilidad. Las predicciones ayudan a priorizar "
    "desviaciones y revisar el impacto económico de las decisiones operativas."
)
add_numbered(
    doc,
    [
        "Abrir /auth/finanzas/dashboard/. Revise las tarjetas, tendencias y alertas del periodo seleccionado.",
        "Compare costo real contra costo planificado y presupuesto. Una desviación debe tener periodo y centro de responsabilidad.",
        "Revise las métricas de flujo, cuentas por cobrar y cuentas por pagar, y confirme que las fechas estén completas.",
        "Use exportación Excel o PDF solo después de validar filtros, periodo y unidad monetaria.",
    ],
)

doc.add_heading("Fase 7 Operación del ciclo predictivo", level=2)
doc.add_paragraph(
    "La séptima fase integra las anteriores en una ejecución repetible. El ciclo calcula predicciones, las registra con versión, "
    "evalúa periodos vencidos cuando existe un resultado real y genera un resumen de salud y deriva del sistema."
)
doc.add_paragraph("Ejecutar el ciclo completo:")
for command in [
    "venv\\Scripts\\python.exe manage.py ejecutar_ciclo_predictivo --usuario admin_ml --json",
    "venv\\Scripts\\python.exe manage.py ejecutar_ciclo_predictivo --usuario admin_ml --sin-evaluar --json",
]:
    paragraph = doc.add_paragraph(style="Intense Quote")
    run = paragraph.add_run(command)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(9.5)
add_bullets(
    doc,
    [
        "La primera ejecución registra las predicciones y evalúa resultados vencidos.",
        "La opción --sin-evaluar sirve para revisar un ciclo de registro sin alterar la evaluación histórica.",
        "Ejecutar dos veces el mismo ciclo debe conservar una sola predicción por clave de periodo, departamento y métrica; esto comprueba la idempotencia.",
        "El resumen reporta cantidad registrada, evaluaciones, pendientes, MAPE y posible deriva.",
    ],
)

doc.add_heading("Simulación controlada", level=2)
doc.add_paragraph(
    "Para preparar datos de prueba se puede generar una serie semanal local. Diez periodos son suficientes para una primera "
    "revisión; en la validación de la Fase 7 se utilizaron doce periodos para aumentar la cobertura temporal."
)
for command in [
    "venv\\Scripts\\python.exe manage.py simular_semana_ml --periods 10",
    "venv\\Scripts\\python.exe manage.py ejecutar_ciclo_predictivo --usuario admin_ml --json",
]:
    paragraph = doc.add_paragraph(style="Intense Quote")
    run = paragraph.add_run(command)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(9.5)
doc.add_paragraph(
    "La simulación se debe usar solo en desarrollo. Antes de ejecutarla, verificar que el usuario entienda que reemplaza "
    "los registros marcados como simulación; no debe ejecutarse sobre datos productivos."
)

doc.add_heading("Cómo interpretar el resultado", level=2)
add_table(
    doc,
    ["Estado", "Interpretación", "Acción del usuario"],
    [
        ("Verde", "Datos suficientes y comportamiento estable", "Continuar operación y revisar la tendencia"),
        ("Amarillo", "Predicción pendiente, poca muestra o calidad media", "Validar datos y esperar resultado real"),
        ("Rojo", "Error alto, deriva o riesgo operativo", "Investigar causa y registrar acción correctiva"),
    ],
    widths=[0.9, 3.0, 2.65],
)
doc.add_paragraph(
    "El estado amarillo no significa que el proceso haya fallado. En particular, las predicciones futuras permanecen "
    "pendientes hasta que se capture el resultado real del periodo proyectado."
)

doc.add_heading("Flujo técnico resumido", level=2)
doc.add_paragraph(
    "Datos operativos → funciones predictivas por departamento → RegistroPrediccion → dashboards y alertas → resultado real → "
    "error y deriva → nueva versión o ajuste del modelo",
    style="Subtitle",
)
doc.add_paragraph(
    "La separación entre cálculo, registro, visualización y evaluación permite auditar cada predicción y evita que un cambio "
    "visual altere la lógica de negocio."
)

doc.add_heading("Lista final de comprobación", level=2)
add_bullets(
    doc,
    [
        "La aplicación abre con la base local y no intenta conectarse a Supabase.",
        "Django check, migraciones y servidor terminan sin errores.",
        "Cada módulo muestra datos históricos y una proyección comprensible.",
        "Las métricas pueden seleccionarse y ocultarse para comparar tendencias.",
        "El ciclo predictivo registra las cinco áreas sin duplicar predicciones.",
        "Los resultados reales permiten evaluar MAPE y detectar deriva.",
        "Las acciones recomendadas tienen responsable o siguiente paso identificable.",
    ],
)

doc.add_heading("Pendientes para una puesta en producción", level=2)
add_bullets(
    doc,
    [
        "Programar el comando de ciclo en el planificador del servidor.",
        "Incorporar pruebas automatizadas para datos, modelos, permisos y vistas.",
        "Definir respaldos y migrar de SQLite a PostgreSQL cuando el entorno productivo esté preparado.",
        "Capturar más eventos reales de mantenimiento para mejorar la calidad de sus predicciones.",
    ],
)

doc.save(OUTPUT)
print(OUTPUT)
