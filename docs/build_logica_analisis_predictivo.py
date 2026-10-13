from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Logica_del_analisis_predictivo_actualizado.docx"


def set_font(run, name="Aptos", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def borders(cell, color="D9D9D9"):
    props = cell._tc.get_or_add_tcPr()
    node = props.first_child_found_in("w:tcBorders")
    if node is None:
        node = OxmlElement("w:tcBorders")
        props.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn("w:" + edge)
        item = node.find(tag)
        if item is None:
            item = OxmlElement("w:" + edge)
            node.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:space"), "0")
        item.set(qn("w:color"), color)


def cell_margins(cell):
    props = cell._tc.get_or_add_tcPr()
    node = props.first_child_found_in("w:tcMar")
    if node is None:
        node = OxmlElement("w:tcMar")
        props.append(node)
    for name in ("top", "start", "bottom", "end"):
        item = OxmlElement("w:" + name)
        item.set(qn("w:w"), "120")
        item.set(qn("w:type"), "dxa")
        node.append(item)


def set_cell(cell, value, bold=False, color="000000", size=9.1):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(str(value))
    set_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)
    borders(cell)


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    props.append(header)


def table(doc, headers, rows, widths, size=9.1):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c, "17365D")
        set_cell(c, header, bold=True, color="FFFFFF", size=size)
        c.width = Inches(widths[i])
    repeat_header(t.rows[0])
    for row_i, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            if row_i % 2:
                shade(cells[i], "F2F6FA")
            set_cell(cells[i], value, size=size)
            cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def formula(doc, text):
    p = doc.add_paragraph(style="Intense Quote")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_font(r, name="Cambria Math", size=11, color="17365D")


def code(doc, text):
    p = doc.add_paragraph(style="Intense Quote")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, name="Consolas", size=9.5, color="17365D")


def page_field(p):
    r = p.add_run("Página ")
    set_font(r, size=8, color="5E6E82")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, instr, end])


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)
styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.8)
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size in (("Title", 25), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
    s = styles[name]
    s.font.name = "Aptos Display"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fr = footer.add_run("Sistema JulMarp ERP · Lógica del análisis predictivo · ")
set_font(fr, size=8, color="5E6E82")
page_field(footer)

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(4)
title.add_run("Lógica del análisis predictivo")
sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(16)
sr = sub.add_run("Explicación del pronóstico por segmento para mi exposición de tesis")
set_font(sr, size=11.5, color="52657E")

doc.add_paragraph(
    "En este documento explico, desde mi propio proyecto, cómo funciona el análisis predictivo que integré en el ERP. "
    "Presento la estructura de los datos, la base matemática, el pronóstico por segmento y el cálculo completo que me permite "
    "obtener una predicción interpretable. Mi objetivo es demostrar de dónde sale cada número y distinguir los datos simulados "
    "de los resultados que posteriormente provengan de la operación real."
)

doc.add_heading("1. Qué problema resuelvo con el análisis predictivo", level=1)
doc.add_paragraph(
    "En mi ERP no solamente registro lo que ya ocurrió. También utilizo el historial operativo para estimar qué puede ocurrir "
    "en las siguientes semanas. Esto me permite anticipar una disminución de producción, identificar el segmento que concentra "
    "el riesgo y convertir el resultado en una acción de seguimiento."
)
doc.add_paragraph(
    "El pronóstico no sustituye la decisión del responsable de producción. Es una señal cuantitativa que me ayuda a revisar "
    "capacidad, programación, materiales, turnos y causas de variación antes de que el problema sea mayor."
)

doc.add_heading("2. Conceptos que utilizo", level=1)
table(doc, ["Concepto", "Cómo lo interpreto en mi sistema"], [
    ("Dato histórico", "Producción real de una semana ya terminada."),
    ("Segmento", "Grupo de órdenes que comparte línea, producto o turno."),
    ("Muestra", "Una semana con datos disponibles para el segmento."),
    ("Actual", "Último total semanal registrado para ese segmento."),
    ("Proyección", "Valor estimado para el final del horizonte de cuatro semanas."),
    ("Cambio", "Diferencia entre proyección y valor actual."),
    ("Confianza", "Señal construida con muestras y error histórico."),
    ("Estado", "Clasificación operativa verde, amarillo o rojo."),
], [1.55, 5.0])

doc.add_heading("3. Estructura general del análisis", level=1)
doc.add_paragraph(
    "La estructura que implementé separa los datos, el cálculo, la evaluación y la presentación. Así puedo explicar cada resultado "
    "y modificar la interfaz sin alterar la fórmula matemática."
)
table(doc, ["Capa", "Qué hace", "Resultado"], [
    ("Fuente operativa", "Lee órdenes de fabricación completadas.", "Datos reales o simulados."),
    ("Agrupación", "Organiza órdenes por semana y segmento.", "Serie histórica por grupo."),
    ("Modelo", "Calcula la tendencia lineal.", "Proyección de cuatro semanas."),
    ("Evaluación", "Compara pronóstico histórico contra resultado real.", "Error y confianza."),
    ("Regla de estado", "Compara proyección contra valor actual.", "Verde, amarillo o rojo."),
    ("Presentación", "Muestra tabla, cambio y recomendación.", "Decisión entendible."),
], [1.35, 3.25, 1.95])
formula(doc, "Órdenes completadas → agrupación semanal → tendencia → proyección → evaluación → estado → acción")

doc.add_heading("4. De dónde salen los datos", level=1)
doc.add_paragraph(
    "Para producción utilizo las órdenes de fabricación con estado completado. De cada orden tomo la fecha de término, la cantidad "
    "producida, la línea, el producto asociado a la BOM y el turno. Si existen varias órdenes del mismo segmento en una semana, "
    "primero sumo sus cantidades."
)
table(doc, ["Campo de la orden", "Uso en el pronóstico"], [
    ("fecha_fin_real", "Define la semana histórica."),
    ("cantidad_producida", "Forma el total semanal que se pronostica."),
    ("linea_produccion", "Forma los segmentos por línea."),
    ("bom.producto", "Forma los segmentos por producto."),
    ("turno", "Forma los segmentos por turno."),
    ("estado completada", "Evita utilizar órdenes todavía no terminadas."),
], [2.0, 4.55])

doc.add_heading("5. Qué significa pronóstico por segmento", level=1)
doc.add_paragraph(
    "El pronóstico por segmento significa que no analizo toda la producción como una sola cifra. La separo en grupos para saber "
    "dónde se concentra la variación. Actualmente observo tres dimensiones: línea, producto y turno."
)
bullets(doc, [
    "Por línea: me permite identificar si una línea pierde capacidad o aumenta su producción.",
    "Por producto: me permite observar si la variación está relacionada con el producto fabricado.",
    "Por turno: me permite revisar si existe una diferencia entre horarios de trabajo.",
])
doc.add_paragraph(
    "Las dimensiones se calculan por separado sobre las mismas órdenes. Por eso una misma tendencia puede aparecer en una línea, "
    "en su producto asociado y en su turno asociado. No significa que se haya duplicado la producción; significa que observo el mismo "
    "comportamiento desde tres perspectivas operativas."
)

doc.add_heading("6. Paso uno: convertir órdenes en una serie semanal", level=1)
doc.add_paragraph(
    "Para cada segmento creo pares de datos. La variable x representa la posición temporal de la semana y la variable y representa "
    "la producción total de esa semana. La variable x no es una cantidad de piezas ni un promedio."
)
table(doc, ["Variable", "Significado", "Ejemplo"], [
    ("x", "Posición temporal de la semana.", "0, 1, 2, ..., 9"),
    ("y", "Producción total semanal del segmento.", "920, 882, ..., 658"),
], [0.75, 3.65, 2.15])
doc.add_paragraph(
    "En el ejemplo de la Línea ML-02 tengo diez semanas con las siguientes cantidades producidas. Estos datos corresponden a la "
    "simulación segmentada utilizada para comprobar la lógica del sistema."
)
table(doc, ["Semana", "x", "Producción y"], [
    ("1", "0", "920"), ("2", "1", "882"), ("3", "2", "872"), ("4", "3", "824"),
    ("5", "4", "813"), ("6", "5", "772"), ("7", "6", "762"), ("8", "7", "718"),
    ("9", "8", "700"), ("10", "9", "658"),
], [1.2, 1.2, 2.0])

doc.add_heading("7. Paso dos: calcular los promedios", level=1)
doc.add_paragraph(
    "Para calcular la recta de tendencia, primero obtengo el promedio de las posiciones x y el promedio de las producciones y. "
    "El promedio de x se obtiene sumando las posiciones y dividiendo entre las diez semanas."
)
formula(doc, "promedio_x = (0 + 1 + 2 + 3 + 4 + 5 + 6 + 7 + 8 + 9) / 10 = 45 / 10 = 4.5")
formula(doc, "promedio_y = (920 + 882 + 872 + 824 + 813 + 772 + 762 + 718 + 700 + 658) / 10 = 7921 / 10 = 792.1")
doc.add_paragraph(
    "El valor 4.5 no representa piezas. Es el punto medio temporal de las diez semanas. El valor 792.1 sí representa el promedio "
    "de piezas producidas en esas semanas."
)

doc.add_heading("8. Paso tres: calcular la pendiente", level=1)
doc.add_paragraph(
    "La pendiente me indica cuánto cambia la producción cuando avanzo una semana. La calculo mediante mínimos cuadrados ordinarios, "
    "buscando la recta que se ajusta mejor a todos los puntos históricos."
)
formula(doc, "b = Σ[(x - promedio_x)(y - promedio_y)] / Σ[(x - promedio_x)²]")
table(doc, ["x", "y", "(x − 4.5)(y − 792.1)", "(x − 4.5)²"], [
    ("0", "920", "-575.55", "20.25"), ("1", "882", "-314.65", "12.25"),
    ("2", "872", "-199.75", "6.25"), ("3", "824", "-47.85", "2.25"),
    ("4", "813", "-10.45", "0.25"), ("5", "772", "-10.05", "0.25"),
    ("6", "762", "-45.15", "2.25"), ("7", "718", "-185.25", "6.25"),
    ("8", "700", "-322.35", "12.25"), ("9", "658", "-603.45", "20.25"),
    ("Total", "", "-2314.50", "82.50"),
], [0.7, 0.9, 3.05, 1.6], size=8.7)
formula(doc, "b = -2314.50 / 82.50 = -28.0545 piezas por semana")
doc.add_paragraph(
    "La pendiente negativa indica que, en este segmento, la producción disminuye aproximadamente 28.05 piezas por semana. "
    "Los valores -2314.50 y 82.50 salen de sumar las operaciones de las diez filas; no provienen de una fuente externa."
)

doc.add_heading("9. Paso cuatro: calcular el intercepto y la ecuación", level=1)
doc.add_paragraph("Después de obtener la pendiente, calculo el intercepto a para construir la ecuación completa de la recta.")
formula(doc, "a = promedio_y - b × promedio_x")
formula(doc, "a = 792.1 - (-28.0545 × 4.5) = 918.3455")
formula(doc, "ŷ = a + b × x")
formula(doc, "ŷ = 918.3455 - 28.0545 × x")

doc.add_heading("10. Paso cinco: proyectar cuatro semanas", level=1)
table(doc, ["Semana futura", "x", "Cálculo", "Proyección"], [
    ("1", "10", "918.3455 − 28.0545 × 10", "637.80 piezas"),
    ("2", "11", "918.3455 − 28.0545 × 11", "609.75 piezas"),
    ("3", "12", "918.3455 − 28.0545 × 12", "581.69 piezas"),
    ("4", "13", "918.3455 − 28.0545 × 13", "553.64 piezas"),
], [1.15, 0.75, 3.35, 1.1], size=8.8)
doc.add_paragraph(
    "El sistema utiliza un horizonte de cuatro semanas. Por eso el valor que aparece como proyección final es 553.64 piezas, "
    "correspondiente a la cuarta semana futura y no solamente a la semana siguiente."
)

doc.add_heading("11. Paso seis: calcular el cambio y el estado", level=1)
formula(doc, "cambio = proyección final - valor actual")
formula(doc, "cambio = 553.64 - 658 = -104.36 piezas")
formula(doc, "porcentaje de cambio = (-104.36 / 658) × 100 = -15.86%")
doc.add_paragraph(
    "Como la proyección final cae más del 10% con respecto al valor actual, el sistema clasifica el segmento como rojo. La clasificación "
    "no se basa solo en que el número sea negativo; se basa en la magnitud relativa de la disminución."
)
table(doc, ["Regla", "Estado", "Interpretación"], [
    ("Proyección menor al 90% del actual", "Rojo", "Riesgo de disminución importante."),
    ("Proyección menor al 97% del actual", "Amarillo", "Disminución que requiere seguimiento."),
    ("Proyección igual o mayor al 97%", "Verde", "Comportamiento estable o favorable."),
], [2.7, 1.0, 2.8])

doc.add_heading("12. Paso siete: calcular la confianza", level=1)
doc.add_paragraph(
    "También calculo una señal de confianza. Para ello considero la cantidad de muestras y el error observado cuando pruebo la tendencia "
    "con periodos históricos anteriores. La confianza no significa certeza; significa que la calidad del historial permite interpretar el "
    "resultado con mayor o menor seguridad."
)
table(doc, ["Condición", "Etiqueta"], [
    ("Menos de 4 muestras", "Baja"),
    ("Al menos 8 muestras y error controlado", "Alta"),
    ("Datos intermedios o error moderado", "Media"),
    ("Error elevado o comportamiento inestable", "Baja"),
], [4.4, 2.1])
doc.add_paragraph("En el ejemplo tengo 10 muestras y una tendencia consistente, por lo que la interfaz muestra confianza alta.")

doc.add_heading("13. Resultado de los nueve segmentos simulados", level=1)
doc.add_paragraph(
    "Para validar visualmente el funcionamiento generé tres combinaciones de línea, producto y turno durante diez semanas completas. "
    "La misma tendencia se puede observar en las tres dimensiones porque cada escenario está asociado a una línea, un producto y un turno."
)
table(doc, ["Dimensión", "Segmento", "Actual", "Proyección", "Cambio", "Estado"], [
    ("Línea", "Línea ML-01", "943", "1084.34", "+141.34", "Verde"),
    ("Línea", "Línea ML-02", "658", "553.64", "-104.36", "Rojo"),
    ("Línea", "Línea ML-03", "772", "763.72", "-8.28", "Verde"),
    ("Producto", "Producto Alfa ML", "943", "1084.34", "+141.34", "Verde"),
    ("Producto", "Producto Beta ML", "658", "553.64", "-104.36", "Rojo"),
    ("Producto", "Producto Gamma ML", "772", "763.72", "-8.28", "Verde"),
    ("Turno", "Matutino", "943", "1084.34", "+141.34", "Verde"),
    ("Turno", "Vespertino", "658", "553.64", "-104.36", "Rojo"),
    ("Turno", "Nocturno", "772", "763.72", "-8.28", "Verde"),
], [1.0, 1.75, 0.85, 1.05, 0.9, 0.9], size=8.5)

doc.add_heading("14. Cómo explico el flujo en mi exposición", level=1)
numbered(doc, [
    "Primero explico que utilizo órdenes de fabricación completadas y no órdenes todavía abiertas.",
    "Después indico que las órdenes se agrupan por semana y por segmento: línea, producto o turno.",
    "Aclaro que x representa la posición temporal y y representa el total de piezas producido.",
    "Presento la regresión lineal como una recta que resume la tendencia histórica.",
    "Muestro cómo la pendiente indica el aumento o la disminución promedio por semana.",
    "Sustituyo posiciones futuras en la ecuación para obtener cuatro proyecciones.",
    "Comparo la última producción real contra la última proyección para calcular el cambio.",
    "Finalmente explico el estado, la confianza y la acción que debe revisar el responsable operativo.",
])
doc.add_paragraph(
    "Una forma breve de expresarlo durante mi exposición es: “Tomo la producción histórica de cada segmento, la organizo por semana, "
    "calculo una tendencia lineal y la extiendo cuatro semanas. Después comparo la proyección con el último resultado real para saber "
    "si el segmento se mantiene estable, requiere seguimiento o presenta un riesgo importante”."
)

doc.add_heading("15. Supuestos y límites que declaro en la tesis", level=1)
bullets(doc, [
    "La regresión lineal supone que la tendencia histórica puede continuar durante el horizonte analizado.",
    "El modelo no demuestra causalidad; identifica una relación temporal entre semanas y producción.",
    "Una falla, falta de material, orden urgente o cambio de capacidad puede separar el resultado real de la proyección.",
    "La simulación segmentada es sintética y sirve para demostrar la funcionalidad; no representa producción real de la empresa.",
    "En una etapa posterior puedo agregar demanda, inventario, paros, calidad, capacidad y estacionalidad.",
    "La confianza debe actualizarse cuando se incorporen resultados reales y se mida el error de manera continua.",
])

doc.add_heading("16. Cómo reproduzco la demostración", level=1)
doc.add_paragraph("Para reproducir el escenario de tesis utilizo diez semanas completas y después ejecuto el ciclo predictivo:")
code(doc, "venv\\Scripts\\python.exe manage.py simular_semana_ml --periods 10")
code(doc, "venv\\Scripts\\python.exe manage.py ejecutar_ciclo_predictivo --usuario admin_ml --json")
doc.add_paragraph(
    "La simulación queda identificada internamente como sintética y puede repetirse sin mezclarse con datos reales. La segunda orden "
    "actualiza el registro predictivo sin crear duplicados, por lo que también puedo comprobar que el ciclo es repetible."
)

doc.add_heading("17. Cómo obtengo el pronóstico de producción", level=1)
doc.add_paragraph(
    "Además del pronóstico por segmento, en mi sistema calculo un pronóstico global de producción. Esta vista no separa inicialmente "
    "la producción por línea, producto o turno; toma los reportes KPI semanales y analiza el comportamiento general de la operación. "
    "Por eso en la pantalla puedo observar la producción esperada, las métricas principales, el riesgo y la gráfica histórica con su proyección."
)

doc.add_heading("17.1 Selección del historial", level=2)
doc.add_paragraph(
    "Primero recupero los reportes KPI de producción ordenados por fecha de cierre. Para cada semana conservo el reporte más reciente, "
    "con el fin de evitar que una misma semana aparezca repetida. El sistema puede conservar hasta 104 semanas, pero para que el pronóstico "
    "sea confiable necesito por lo menos cuatro periodos y en la simulación utilizo doce muestras."
)
table(doc, ["Fuente", "Qué obtengo"], [
    ("Reporte KPI semanal", "Una observación histórica por semana."),
    ("Detalle del reporte", "Producción total y datos operativos complementarios."),
    ("Fecha de cierre", "Orden temporal y etiquetas de la gráfica."),
    ("Horizonte seleccionado", "Cantidad de semanas futuras que se proyectan."),
], [2.25, 4.3])

doc.add_heading("17.2 Construcción de la serie de producción", level=2)
doc.add_paragraph(
    "Para la métrica de producción total tomo el valor de unidades totales registrado en el detalle de cada reporte KPI. Con esos "
    "valores formo la serie histórica y utilizo la fecha de cierre para ubicar cada observación en el tiempo. En lugar de suponer que "
    "todas las semanas están perfectamente separadas, calculo la posición temporal como la cantidad de semanas transcurridas desde el primer reporte."
)
formula(doc, "x = (fecha de cierre actual - fecha del primer reporte) / 7 días")
formula(doc, "y = unidades totales producidas en el reporte KPI")
doc.add_paragraph(
    "Así, x representa el tiempo transcurrido y y representa la producción total. Si existe más de una orden dentro de un mismo periodo, "
    "el reporte KPI ya concentra ese comportamiento en el valor semanal que utilizo para el análisis."
)

doc.add_heading("17.3 Aplicación de la tendencia lineal", level=2)
doc.add_paragraph(
    "Sobre la serie de producción aplico la misma base matemática de regresión lineal explicada para el pronóstico por segmento. Calculo "
    "la pendiente con todos los reportes disponibles, obtengo el intercepto y construyo una ecuación que representa el comportamiento global."
)
formula(doc, "ŷ = a + b × x")
doc.add_paragraph(
    "Cuando el sistema tiene menos de tres valores, utiliza un promedio ponderado como respaldo. Cuando tiene cuatro o más muestras, la "
    "proyección se considera suficientemente confiable para mostrarse como resultado del modelo. Con doce muestras, la vista puede mostrar "
    "confianza alta si el error histórico también se mantiene controlado."
)

doc.add_heading("17.4 Proyección del horizonte futuro", level=2)
doc.add_paragraph(
    "Si selecciono cuatro semanas, calculo cuatro valores futuros sustituyendo en la ecuación las posiciones posteriores al último reporte. "
    "La producción esperada que aparece en la tarjeta resumen corresponde al último valor de ese horizonte."
)
table(doc, ["Elemento de la pantalla", "Cómo lo obtengo"], [
    ("12 muestras", "Número de reportes KPI semanales utilizados."),
    ("Próximas 4 semanas", "Horizonte elegido para calcular cuatro valores futuros."),
    ("Producción esperada", "Último valor de la serie futura de producción total."),
    ("Histórico", "Valores reales de los reportes KPI."),
    ("Pronóstico", "Valores calculados para las semanas posteriores."),
], [2.2, 4.35])

doc.add_heading("17.5 Rango de incertidumbre", level=2)
doc.add_paragraph(
    "Para no mostrar la proyección como un número exacto, calculo un margen a partir del error histórico. Si existen validaciones retrospectivas, "
    "utilizo el error absoluto medio. Si todavía no hay suficientes validaciones, utilizo como referencia la variación promedio entre periodos consecutivos."
)
formula(doc, "margen = error histórico × raíz cuadrada del horizonte")
formula(doc, "límite inferior = proyección - margen")
formula(doc, "límite superior = proyección + margen")
doc.add_paragraph(
    "En las métricas expresadas como porcentaje, mantengo los límites entre 0% y 100%. Para producción total evito que el límite inferior "
    "sea negativo. Este rango representa incertidumbre del modelo; no es una garantía ni un compromiso de producción."
)

doc.add_heading("17.6 Evaluación y alertas", level=2)
doc.add_paragraph(
    "Para evaluar la calidad del pronóstico, simulo retrospectivamente el cálculo: tomo una parte de los periodos históricos, proyecto el siguiente "
    "y comparo el valor calculado con el resultado que realmente ocurrió. Con esas diferencias calculo MAE, MAPE y la confianza."
)
table(doc, ["Indicador", "Qué me permite observar"], [
    ("MAE", "Cuántas unidades, puntos o minutos se desvía en promedio."),
    ("MAPE", "Qué porcentaje representa el error respecto al valor real."),
    ("Confianza", "Qué tan consistente es el historial disponible."),
    ("Alerta", "Qué métrica necesita una revisión operativa."),
], [1.3, 5.25])
doc.add_paragraph(
    "Por ejemplo, si la proyección de rendimiento se deteriora o el tiempo de ciclo aumenta de manera importante, el sistema coloca la métrica "
    "en las alertas principales y muestra una acción sugerida. La alerta no se genera solamente por una cifra aislada; se apoya en la tendencia "
    "proyectada y en las reglas del indicador."
)

doc.add_heading("17.7 Diferencia entre ambos pronósticos", level=2)
table(doc, ["Pronóstico global", "Pronóstico por segmento"], [
    ("Utiliza reportes KPI semanales.", "Utiliza órdenes de fabricación completadas."),
    ("Observa la producción total y sus KPI.", "Separa línea, producto y turno."),
    ("Muestra la tarjeta de producción esperada y la gráfica general.", "Muestra actual, proyección y cambio por grupo."),
    ("Ayuda a conocer el estado general de MFG.", "Ayuda a localizar dónde se concentra la variación."),
], [3.25, 3.3])
doc.add_paragraph(
    "En mi exposición puedo explicar que ambos análisis utilizan la misma base de tendencia lineal, pero parten de diferentes niveles de detalle. "
    "El pronóstico global responde cuánto espero producir en general; el pronóstico por segmento responde en qué línea, producto o turno se concentra "
    "la variación."
)

doc.add_heading("17.8 Guion breve para explicar el pronóstico de producción", level=2)
doc.add_paragraph(
    "“Para obtener el pronóstico de producción tomo los reportes KPI cerrados por semana, selecciono una observación por periodo y construyo una serie "
    "histórica de producción total. Después asigno a cada reporte una posición temporal, calculo una regresión lineal y proyecto las siguientes semanas. "
    "La tarjeta de producción esperada muestra el último valor del horizonte elegido. Finalmente comparo el pronóstico con el historial, calculo el rango "
    "de incertidumbre y genero alertas cuando una métrica presenta una tendencia que requiere atención”."
)

doc.save(OUTPUT)
print(OUTPUT)
